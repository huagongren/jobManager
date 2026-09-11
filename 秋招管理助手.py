"""秋招管理助手：本地岗位投递管理桌面软件。"""

import sqlite3
import sys
import re
import os
import shutil
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import webbrowser

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    OPENPYXL_ERROR = None
except ImportError as error:
    Workbook = load_workbook = None
    Alignment = Font = PatternFill = None
    OPENPYXL_ERROR = error


APP_DIR = Path(os.getenv("CAMPUS_JOB_TRACKER_HOME") or (Path.home() / ".campus_job_tracker"))
DB_PATH = APP_DIR / "jobs.db"
ATTACHMENTS_DIR = APP_DIR / "attachments"  # legacy location kept for one-time migration
DOCUMENTS_DIR = APP_DIR / "documents"
DOCUMENT_CATEGORIES = ("简历", "成绩单", "证书", "作品集", "其他")
STATUSES = [
    "待投递", "申请", "已投递", "简历挂", "笔试", "笔试挂", "AI面",
    "一面", "一面挂", "二面", "二面挂", "终面", "终面挂", "Offer", "已拒绝", "已终止",
]
TERMINATED_STATUSES = {"简历挂", "笔试挂", "一面挂", "二面挂", "终面挂", "已拒绝", "已终止"}
CALENDAR_EXCLUDED_STATUSES = TERMINATED_STATUSES | {"Offer"}
ONGOING_STATUSES = {"申请", "已投递", "笔试", "AI面", "一面", "二面", "终面"}
REMINDER_DAYS = 3
EXCEL_COLUMNS = [
    ("company", "公司名称"), ("position", "岗位名称"), ("location", "工作地点"),
    ("salary", "薪资范围"), ("channel", "投递渠道"), ("status", "当前状态"),
    ("applied_date", "投递日期"), ("interview_time", "笔试/面试时间"),
    ("deadline", "截止日期"), ("link", "岗位链接"), ("notes", "备注"),
]
EXCEL_HEADER_ALIASES = {
    "公司": "company", "公司名称": "company", "企业": "company", "company": "company",
    "岗位": "position", "职位": "position", "岗位名称": "position", "position": "position",
    "地点": "location", "城市": "location", "工作地点": "location", "location": "location",
    "薪资": "salary", "薪酬": "salary", "工资": "salary", "薪资范围": "salary", "salary": "salary",
    "渠道": "channel", "来源": "channel", "投递渠道": "channel", "channel": "channel",
    "状态": "status", "投递状态": "status", "当前状态": "status", "status": "status",
    "投递日期": "applied_date", "投递时间": "applied_date", "申请日期": "applied_date", "applied_date": "applied_date",
    "面试时间": "interview_time", "笔试时间": "interview_time", "笔试/面试时间": "interview_time", "interview_time": "interview_time",
    "截止日期": "deadline", "截止时间": "deadline", "deadline": "deadline",
    "链接": "link", "官网": "link", "岗位链接": "link", "link": "link",
    "备注": "notes", "notes": "notes",
}
STATUS_COLORS = {
    "待投递": "#64748b", "申请": "#3b82f6", "已投递": "#3b82f6", "简历挂": "#94a3b8", "笔试": "#8b5cf6", "AI面": "#a855f7",
    "笔试挂": "#94a3b8", "一面": "#f59e0b", "一面挂": "#94a3b8", "二面": "#f97316", "二面挂": "#94a3b8", "终面": "#ef4444", "终面挂": "#94a3b8",
    "Offer": "#10b981", "已拒绝": "#94a3b8", "已终止": "#475569",
}
STATUS_TINTS = {
    "待投递": "#f1f5f9", "申请": "#eff6ff", "已投递": "#eff6ff", "简历挂": "#f8fafc",
    "笔试": "#f5f3ff", "笔试挂": "#f8fafc", "AI面": "#faf5ff", "一面": "#fffbeb", "一面挂": "#f8fafc",
    "二面": "#fff7ed", "二面挂": "#f8fafc", "终面": "#fef2f2", "终面挂": "#f8fafc",
    "Offer": "#ecfdf5", "已拒绝": "#f8fafc", "已终止": "#f1f5f9",
}


class ScrollableCard(ttk.Frame):
    """A vertically scrollable card so every form field remains reachable."""
    def __init__(self, parent):
        super().__init__(parent, style="Card.TFrame")
        self.canvas = tk.Canvas(self, bg="white", highlightthickness=0, borderwidth=0)
        self.scrollbar = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.content = ttk.Frame(self.canvas, style="Card.TFrame", padding=18)
        self.window = self.canvas.create_window((0, 0), window=self.content, anchor="nw")
        self.canvas.configure(yscrollcommand=self.scrollbar.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        self.scrollbar.pack(side="right", fill="y")
        self.content.bind("<Configure>", self._update_scroll_region)
        self.canvas.bind("<Configure>", self._fit_content_width)
        # The form contains many child widgets. Listen at the application level
        # and scroll only while the cursor is over this card.
        self.bind_all("<MouseWheel>", self._mousewheel, add="+")

    def _update_scroll_region(self, _event=None):
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def _fit_content_width(self, event):
        self.canvas.itemconfigure(self.window, width=event.width)

    def _mousewheel(self, event):
        hovered = self.winfo_containing(self.winfo_pointerx(), self.winfo_pointery())
        if hovered is None:
            return
        content_path = str(self.content)
        if hovered == self.canvas or str(hovered).startswith(content_path):
            direction = -1 if event.delta > 0 else 1
            steps = max(1, abs(event.delta) // 120)
            self.canvas.yview_scroll(direction * steps, "units")
            return "break"


class JobTracker(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("秋招管理助手")
        self.geometry("1220x850")
        self.minsize(1000, 700)
        self.configure(bg="#f6f8fc")
        self.selected_id = None
        self._setup_database()
        self._setup_style()
        self._build_ui()
        self.refresh()
        self.protocol("WM_DELETE_WINDOW", self.close)
        self.after(700, self.check_deadline_reminders)

    def _setup_database(self):
        APP_DIR.mkdir(parents=True, exist_ok=True)
        ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
        DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(DB_PATH)
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys = ON")
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS jobs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company TEXT NOT NULL,
                position TEXT NOT NULL,
                location TEXT,
                salary TEXT,
                channel TEXT,
                status TEXT NOT NULL DEFAULT '待投递',
                deadline DATE,
                applied_date DATE,
                interview_time DATETIME,
                link TEXT,
                notes TEXT,
                updated_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS attachments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                job_id INTEGER NOT NULL,
                category TEXT NOT NULL DEFAULT '其他',
                original_name TEXT NOT NULL,
                stored_path TEXT NOT NULL,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(job_id) REFERENCES jobs(id) ON DELETE CASCADE
            )
        """)
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS documents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                category TEXT NOT NULL DEFAULT '其他',
                original_name TEXT NOT NULL,
                stored_path TEXT NOT NULL,
                source_attachment_id INTEGER UNIQUE,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        self._ensure_columns()
        self._normalize_existing_dates()
        self._migrate_legacy_attachments()
        self.conn.commit()

    def _ensure_columns(self):
        """Keep databases created by earlier app versions compatible."""
        existing = {row[1] for row in self.conn.execute("PRAGMA table_info(jobs)")}
        for name, definition in (("salary", "TEXT"), ("channel", "TEXT"), ("interview_time", "TEXT")):
            if name not in existing:
                self.conn.execute(f"ALTER TABLE jobs ADD COLUMN {name} {definition}")

    def _normalize_existing_dates(self):
        for row in self.conn.execute("SELECT id, applied_date, deadline, interview_time FROM jobs").fetchall():
            updates = {}
            for field, with_time in (("applied_date", False), ("deadline", False), ("interview_time", True)):
                raw = row[field] or ""
                if not raw:
                    continue
                try:
                    normalized = self._normalize_date_value(raw, with_time=with_time)
                except ValueError:
                    normalized = raw
                if normalized != raw:
                    updates[field] = normalized
            if updates:
                assignments = ", ".join(f"{field} = ?" for field in updates)
                self.conn.execute(
                    f"UPDATE jobs SET {assignments}, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                    tuple(updates.values()) + (row["id"],),
                )

    def _migrate_legacy_attachments(self):
        """Move attachments created by the previous job-detail UI into Documents."""
        table = self.conn.execute(
            "SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'attachments'"
        ).fetchone()
        if not table:
            return
        rows = self.conn.execute("SELECT * FROM attachments ORDER BY id").fetchall()
        for row in rows:
            if self.conn.execute(
                "SELECT 1 FROM documents WHERE source_attachment_id = ?", (row["id"],)
            ).fetchone():
                continue
            source = (ATTACHMENTS_DIR / row["stored_path"]).resolve()
            if not source.is_file():
                continue
            original_name = Path(row["original_name"] or source.name).name
            stored_name = f"legacy_{uuid.uuid4().hex}_{original_name}"
            target = DOCUMENTS_DIR / stored_name
            try:
                shutil.copy2(source, target)
            except OSError:
                continue
            self.conn.execute(
                "INSERT INTO documents (category, original_name, stored_path, source_attachment_id) VALUES (?, ?, ?, ?)",
                (row["category"] or "其他", original_name, stored_name, row["id"]),
            )

    def _setup_style(self):
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("TFrame", background="#f4f7fb")
        style.configure("Card.TFrame", background="white")
        style.configure("CardTitle.TLabel", background="white", foreground="#15213b", font=("Microsoft YaHei UI", 13, "bold"))
        style.configure("Section.TLabel", background="white", foreground="#1e3a8a", font=("Microsoft YaHei UI", 10, "bold"))
        style.configure("Muted.TLabel", background="white", foreground="#718096", font=("Microsoft YaHei UI", 9))
        style.configure("Field.TLabel", background="white", foreground="#475569", font=("Microsoft YaHei UI", 9, "bold"))
        style.configure("Treeview", rowheight=38, font=("Microsoft YaHei UI", 10), background="white", fieldbackground="white", foreground="#24324a", borderwidth=0)
        style.configure("Treeview.Heading", font=("Microsoft YaHei UI", 10, "bold"), background="#f0f4ff", foreground="#34435e", relief="flat", padding=(10, 9))
        style.map("Treeview", background=[("selected", "#dbeafe")], foreground=[("selected", "#1e3a8a")])
        style.configure("Accent.TButton", background="#2563eb", foreground="white", font=("Microsoft YaHei UI", 10, "bold"), borderwidth=0, padding=(14, 8))
        style.map("Accent.TButton", background=[("active", "#1d4ed8")])
        style.configure("Soft.TButton", background="#eff6ff", foreground="#1d4ed8", font=("Microsoft YaHei UI", 10, "bold"), borderwidth=0, padding=(12, 7))
        style.map("Soft.TButton", background=[("active", "#dbeafe")])
        style.configure("Danger.TButton", background="#fff1f2", foreground="#be123c", font=("Microsoft YaHei UI", 10, "bold"), borderwidth=0, padding=(12, 7))
        style.map("Danger.TButton", background=[("active", "#ffe4e6")])
        style.configure("TButton", font=("Microsoft YaHei UI", 10), padding=(10, 7))
        style.configure("TEntry", padding=8, font=("Microsoft YaHei UI", 10))
        style.configure("TCombobox", padding=7, font=("Microsoft YaHei UI", 10))
        style.configure("Vertical.TScrollbar", troughcolor="#f8fafc", bordercolor="#f8fafc", background="#cbd5e1", arrowcolor="#64748b")
        style.configure("TNotebook", background="#f4f7fb", borderwidth=0)
        style.configure("TNotebook.Tab", background="#e8edf6", foreground="#54627a", padding=(20, 10), font=("Microsoft YaHei UI", 10, "bold"))
        style.map("TNotebook.Tab", background=[("selected", "#111b34")], foreground=[("selected", "white")])

    def _build_ui(self):
        container = ttk.Frame(self, padding=(20, 16))
        container.pack(fill="both", expand=True)

        self.pages = ttk.Notebook(container)
        self.pages.pack(fill="both", expand=True)
        home = ttk.Frame(self.pages)
        documents = ttk.Frame(self.pages)
        self.pages.add(home, text="  首页 · 投递看板  ")
        self.pages.add(documents, text="  简历与材料管理  ")
        self._build_home(home)
        self._build_documents_page(documents)

    def _build_home(self, container):

        header = tk.Frame(container, bg="#111b34", height=108)
        header.pack(fill="x")
        header.pack_propagate(False)
        header_text = tk.Frame(header, bg="#111b34")
        header_text.pack(side="left", padx=24, pady=18)
        tk.Label(header_text, text="秋招控制台", bg="#111b34", fg="white", font=("Microsoft YaHei UI", 22, "bold")).pack(anchor="w")
        tk.Label(header_text, text="每一次投递，都是下一次机会的开始", bg="#111b34", fg="#b9c7e5", font=("Microsoft YaHei UI", 10)).pack(anchor="w", pady=(4, 0))
        tk.Label(header, text="2027届 · 求职档案", bg="#1d2b4d", fg="#dbeafe", padx=12, pady=6, font=("Microsoft YaHei UI", 9, "bold")).pack(side="right", padx=(0, 12), pady=32)
        tk.Button(header, text="＋  新增岗位", command=self.new_job, bg="#5eead4", fg="#10233f", activebackground="#99f6e4", activeforeground="#10233f", relief="flat", borderwidth=0, padx=18, pady=9, font=("Microsoft YaHei UI", 10, "bold"), cursor="hand2").pack(side="right", padx=0, pady=27)

        self.cards = ttk.Frame(container)
        self.cards.pack(fill="x", pady=(16, 16))
        self.card_labels = {}
        card_specs = [
            ("岗位总数", "全部", "#eef4ff", "#3565d8", "已建立求职档案"),
            ("待投递", "待投递", "#f2f5f8", "#64748b", "等待你主动出击"),
            ("流程进行中", "进行中", "#f4efff", "#7c3aed", "投递 / 笔试 / 面试"),
            ("拿到 Offer", "Offer", "#ecfdf5", "#059669", "继续保持节奏"),
        ]
        for i, (title, key, background, accent, hint) in enumerate(card_specs):
            card = tk.Frame(self.cards, bg=background, highlightbackground="#e8edf6", highlightthickness=1)
            card.grid(row=0, column=i, sticky="nsew", padx=(0 if i == 0 else 10, 0))
            tk.Frame(card, bg=accent, height=4).pack(fill="x")
            body = tk.Frame(card, bg=background)
            body.pack(fill="both", expand=True, padx=16, pady=11)
            tk.Label(body, text=title, bg=background, fg="#55657d", font=("Microsoft YaHei UI", 9, "bold")).pack(anchor="w")
            label = tk.Label(body, text="0", bg=background, fg="#15213b", font=("Microsoft YaHei UI", 24, "bold"))
            label.pack(anchor="w", pady=(1, 0))
            tk.Label(body, text=hint, bg=background, fg=accent, font=("Microsoft YaHei UI", 8)).pack(anchor="w")
            self.card_labels[key] = label
            self.cards.columnconfigure(i, weight=1)

        content = ttk.Frame(container)
        content.pack(fill="both", expand=True)
        content.columnconfigure(0, weight=3)
        content.columnconfigure(1, weight=2)
        content.rowconfigure(0, weight=1)

        left = ttk.Frame(content, style="Card.TFrame", padding=18)
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 16))
        self._build_list(left)
        right = ScrollableCard(content)
        right.grid(row=0, column=1, sticky="nsew")
        self._build_form(right.content)

    def _build_documents_page(self, parent):
        page = ttk.Frame(parent, padding=(4, 10))
        page.pack(fill="both", expand=True)
        hero = tk.Frame(page, bg="#172554", height=96)
        hero.pack(fill="x")
        hero.pack_propagate(False)
        hero_text = tk.Frame(hero, bg="#172554")
        hero_text.pack(side="left", padx=24, pady=16)
        tk.Label(hero_text, text="简历与材料管理", bg="#172554", fg="white", font=("Microsoft YaHei UI", 19, "bold")).pack(anchor="w")
        tk.Label(hero_text, text="统一保存求职材料，双击记录即可快速访问", bg="#172554", fg="#bfdbfe", font=("Microsoft YaHei UI", 10)).pack(anchor="w", pady=(4, 0))
        tk.Label(hero, text="上传即复制 · 本地管理", bg="#1e3a6a", fg="#dbeafe", padx=12, pady=6, font=("Microsoft YaHei UI", 9, "bold")).pack(side="right", padx=22, pady=29)

        card = ttk.Frame(page, style="Card.TFrame", padding=20)
        card.pack(fill="both", expand=True, pady=(14, 0))
        toolbar = ttk.Frame(card, style="Card.TFrame")
        toolbar.pack(fill="x", pady=(0, 14))
        ttk.Label(toolbar, text="材料清单", style="CardTitle.TLabel").pack(side="left")
        self.document_count_var = tk.StringVar(value="0 个文件")
        ttk.Label(toolbar, textvariable=self.document_count_var, style="Muted.TLabel").pack(side="left", padx=(10, 0))
        self.document_category_var = tk.StringVar(value="其他")
        ttk.Combobox(toolbar, textvariable=self.document_category_var, values=DOCUMENT_CATEGORIES, state="readonly", width=9).pack(side="right")
        ttk.Label(toolbar, text="类别", style="Muted.TLabel").pack(side="right", padx=(0, 6))
        ttk.Button(toolbar, text="上传其他材料", style="Soft.TButton", command=lambda: self.upload_documents(self.document_category_var.get())).pack(side="right", padx=(8, 0))
        ttk.Button(toolbar, text="上传成绩单", style="Soft.TButton", command=lambda: self.upload_documents("成绩单")).pack(side="right", padx=(8, 0))
        ttk.Button(toolbar, text="上传简历", style="Accent.TButton", command=lambda: self.upload_documents("简历")).pack(side="right", padx=(8, 0))

        table_area = ttk.Frame(card, style="Card.TFrame")
        table_area.pack(fill="both", expand=True)
        table_area.columnconfigure(0, weight=1)
        table_area.rowconfigure(0, weight=1)
        self.document_tree = ttk.Treeview(table_area, columns=("category", "name", "created", "path"), show="headings", selectmode="browse")
        headings = {"category": "类别", "name": "文件名", "created": "添加时间", "path": "本地状态"}
        widths = {"category": 100, "name": 420, "created": 160, "path": 120}
        for column in ("category", "name", "created", "path"):
            self.document_tree.heading(column, text=headings[column])
            self.document_tree.column(column, width=widths[column], minwidth=90, anchor=("w" if column == "name" else "center"))
        self.document_tree.grid(row=0, column=0, sticky="nsew")
        scrollbar = ttk.Scrollbar(table_area, orient="vertical", command=self.document_tree.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.document_tree.configure(yscrollcommand=scrollbar.set)
        self.document_tree.bind("<Double-1>", lambda _event: self.open_document())

        action_row = ttk.Frame(card, style="Card.TFrame")
        action_row.pack(fill="x", pady=(12, 0))
        self.document_hint = tk.StringVar(value=f"文件会复制到：{DOCUMENTS_DIR}")
        ttk.Label(action_row, textvariable=self.document_hint, style="Muted.TLabel").pack(side="left")
        ttk.Button(action_row, text="打开文件夹", style="Soft.TButton", command=self.open_documents_folder).pack(side="right")
        ttk.Button(action_row, text="打开选中文件", style="Soft.TButton", command=self.open_document).pack(side="right", padx=6)
        ttk.Button(action_row, text="删除选中文件", style="Danger.TButton", command=self.delete_document).pack(side="right")
        self.refresh_documents()

    def _build_list(self, parent):
        toolbar = ttk.Frame(parent, style="Card.TFrame")
        toolbar.pack(fill="x", pady=(0, 14))
        toolbar_left = ttk.Frame(toolbar, style="Card.TFrame")
        toolbar_left.pack(side="left")
        ttk.Label(toolbar_left, text="投递进度", style="CardTitle.TLabel").pack(anchor="w")
        self.list_count_var = tk.StringVar(value="0 条记录")
        ttk.Label(toolbar_left, textvariable=self.list_count_var, style="Muted.TLabel").pack(anchor="w", pady=(2, 0))
        toolbar_actions = ttk.Frame(toolbar, style="Card.TFrame")
        toolbar_actions.pack(side="left", padx=(18, 0))
        ttk.Button(toolbar_actions, text="导入 Excel", style="Soft.TButton", command=self.import_excel).pack(side="left", padx=(0, 6))
        ttk.Button(toolbar_actions, text="导出 Excel", style="Soft.TButton", command=self.export_excel).pack(side="left", padx=(0, 6))
        ttk.Button(toolbar_actions, text="导出日历", style="Soft.TButton", command=self.export_calendar).pack(side="left", padx=(0, 6))
        ttk.Button(toolbar_actions, text="检查提醒", style="Soft.TButton", command=lambda: self.check_deadline_reminders(show_empty=True)).pack(side="left")
        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", lambda *_: self.load_jobs())
        ttk.Entry(toolbar, textvariable=self.search_var, width=19).pack(side="right")
        ttk.Label(toolbar, text="搜索", style="Muted.TLabel").pack(side="right", padx=(0, 6))
        self.filter_var = tk.StringVar(value="全部状态")
        status_filter = ttk.Combobox(toolbar, textvariable=self.filter_var, values=["全部状态"] + STATUSES, state="readonly", width=10)
        status_filter.pack(side="right", padx=(0, 12))
        status_filter.bind("<<ComboboxSelected>>", lambda _e: self.load_jobs())

        columns = ("company", "position", "location", "salary", "channel", "status", "applied", "interview", "notes", "link")
        table_area = ttk.Frame(parent, style="Card.TFrame")
        table_area.pack(fill="both", expand=True)
        table_area.columnconfigure(0, weight=1)
        table_area.rowconfigure(0, weight=1)
        self.tree = ttk.Treeview(table_area, columns=columns, show="headings", selectmode="browse")
        headings = {"company": "公司", "position": "岗位", "location": "地点", "salary": "薪资范围", "channel": "投递渠道", "status": "当前状态", "applied": "投递日期", "interview": "笔试/面试时间", "notes": "备注", "link": "官网链接"}
        widths = {"company": 120, "position": 150, "location": 90, "salary": 150, "channel": 90, "status": 90, "applied": 100, "interview": 125, "notes": 180, "link": 90}
        for col in columns:
            self.tree.heading(col, text=headings[col])
            self.tree.column(col, width=widths[col], minwidth=70, anchor=("center" if col not in ("company", "position") else "w"))
        for status, tint in STATUS_TINTS.items():
            self.tree.tag_configure(status, background=tint, foreground="#24324a")
        self.tree.grid(row=0, column=0, sticky="nsew")
        scrollbar = ttk.Scrollbar(table_area, orient="vertical", command=self.tree.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        x_scrollbar = ttk.Scrollbar(table_area, orient="horizontal", command=self.tree.xview)
        x_scrollbar.grid(row=1, column=0, sticky="ew")
        self.tree.configure(yscrollcommand=scrollbar.set, xscrollcommand=x_scrollbar.set)
        self.tree.bind("<<TreeviewSelect>>", self.select_job)
        self.tree.bind("<Double-1>", self.select_job)

    def _build_form(self, parent):
        detail_header = ttk.Frame(parent, style="Card.TFrame")
        detail_header.pack(fill="x")
        ttk.Label(detail_header, text="岗位详情", style="CardTitle.TLabel").pack(side="left")
        self.open_button = ttk.Button(detail_header, text="↗ 打开链接", style="Soft.TButton", command=self.open_link, state="disabled")
        self.open_button.pack(side="right")
        ttk.Label(parent, text="选择一条记录即可查看、修改并保存。", style="Muted.TLabel").pack(anchor="w", pady=(3, 14))
        ttk.Separator(parent, orient="horizontal").pack(fill="x", pady=(0, 12))
        ttk.Label(parent, text="录 入 信 息", style="Section.TLabel").pack(anchor="w", pady=(0, 8))
        self.vars = {name: tk.StringVar() for name in ("company", "position", "location", "salary", "channel", "deadline", "applied_date", "interview_time", "link")}
        self.status_var = tk.StringVar(value="待投递")
        fields = [("公司名称 *", "company"), ("岗位名称 *", "position"), ("工作地点", "location"), ("薪资范围", "salary"), ("投递渠道", "channel"), ("当前状态", "status"), ("投递日期（YYYY-MM-DD）", "applied_date"), ("笔试/面试时间（YYYY-MM-DD HH:MM）", "interview_time"), ("截止日期（YYYY-MM-DD）", "deadline"), ("岗位链接", "link")]
        form_grid = ttk.Frame(parent, style="Card.TFrame")
        form_grid.pack(fill="x")
        form_grid.columnconfigure(0, weight=1)
        form_grid.columnconfigure(1, weight=1)
        for index, (title, name) in enumerate(fields):
            field = ttk.Frame(form_grid, style="Card.TFrame")
            field.grid(row=index // 2, column=index % 2, sticky="ew", padx=(0 if index % 2 == 0 else 10, 0), pady=(0, 10))
            ttk.Label(field, text=title, style="Field.TLabel").pack(anchor="w", pady=(0, 4))
            if name == "status":
                ttk.Combobox(field, textvariable=self.status_var, values=STATUSES, state="readonly").pack(fill="x")
            else:
                ttk.Entry(field, textvariable=self.vars[name]).pack(fill="x")
        ttk.Label(parent, text="备注", style="Field.TLabel").pack(anchor="w", pady=(4, 4))
        self.notes = tk.Text(parent, height=5, font=("Microsoft YaHei UI", 10), bg="#f8fafc", fg="#24324a", relief="flat", borderwidth=0, highlightthickness=1, highlightbackground="#dbe3ef", highlightcolor="#93c5fd", padx=9, pady=7)
        self.notes.pack(fill="x")
        button_row = ttk.Frame(parent, style="Card.TFrame")
        button_row.pack(fill="x", pady=(16, 8))
        self.save_button = ttk.Button(button_row, text="保存岗位", style="Accent.TButton", command=self.save_job)
        self.save_button.pack(side="left")
        ttk.Button(button_row, text="重置", style="Soft.TButton", command=self.new_job).pack(side="left", padx=8)
        self.delete_button = ttk.Button(button_row, text="删除", style="Danger.TButton", command=self.delete_job, state="disabled")
        self.delete_button.pack(side="right")

    @staticmethod
    def _normalize_date_value(value, with_time=False):
        """Normalize user/Excel date input to ISO date or ISO datetime text."""
        text = str(value or "").strip()
        if not text:
            return ""
        text = text.replace("T", " ").replace("t", " ")
        parts = text.split(None, 1)
        date_part = parts[0]
        time_part = parts[1].strip() if len(parts) == 2 else ""

        match = re.fullmatch(r"(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})", date_part)
        if not match:
            match = re.fullmatch(r"(20\d{2})年(\d{1,2})月(\d{1,2})日?", date_part)
        if match:
            year, month, day = (int(item) for item in match.groups())
        else:
            match = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})", date_part)
            if not match:
                match = re.fullmatch(r"(\d{1,2})月(\d{1,2})日?", date_part)
            if not match:
                raise ValueError("日期请使用 YYYY-MM-DD 格式，例如 2026-09-11")
            year = date.today().year
            month, day = (int(item) for item in match.groups())

        try:
            parsed_date = date(year, month, day)
        except ValueError as error:
            raise ValueError(f"日期无效：{text}") from error

        if not with_time:
            if time_part:
                raise ValueError("日期字段不应包含时间，请使用 YYYY-MM-DD")
            return parsed_date.isoformat()
        if not time_part:
            return parsed_date.isoformat()

        time_part = time_part.replace("：", ":")
        time_match = re.fullmatch(r"(\d{1,2}):(\d{1,2})(?::\d{1,2})?", time_part)
        if not time_match:
            time_match = re.fullmatch(r"(\d{1,2})点(?:(\d{1,2})分?)?", time_part)
        if not time_match:
            raise ValueError("时间请使用 HH:MM 格式，例如 14:30")
        hour = int(time_match.group(1))
        minute = int(time_match.group(2) or 0)
        if hour > 23 or minute > 59:
            raise ValueError(f"时间无效：{text}")
        return datetime(year, month, day, hour, minute).strftime("%Y-%m-%d %H:%M")

    @classmethod
    def _date_from_value(cls, value):
        if not value:
            return None
        try:
            normalized = cls._normalize_date_value(value, with_time=True)
            return date.fromisoformat(normalized[:10])
        except (TypeError, ValueError):
            return None

    def _normalize_job_dates(self, values):
        normalized = dict(values)
        for field, label, with_time in (
            ("applied_date", "投递日期", False),
            ("interview_time", "笔试/面试时间", True),
            ("deadline", "截止日期", False),
        ):
            try:
                normalized[field] = self._normalize_date_value(values.get(field, ""), with_time=with_time)
            except ValueError as error:
                messagebox.showwarning("日期格式不正确", f"{label}：{error}")
                return None
        return normalized

    @staticmethod
    def _cell_to_text(value, date_only=False):
        if value is None:
            return ""
        if isinstance(value, datetime):
            if date_only:
                return value.date().isoformat()
            return value.strftime("%Y-%m-%d %H:%M")
        if isinstance(value, date):
            return value.isoformat()
        return str(value).strip()

    def _excel_date_value(self, value, with_time=False):
        normalized = self._normalize_date_value(value, with_time=with_time)
        if not normalized:
            return None
        if with_time and len(normalized) > 10:
            return datetime.strptime(normalized, "%Y-%m-%d %H:%M")
        return date.fromisoformat(normalized[:10])

    def _has_openpyxl(self):
        if Workbook is not None and load_workbook is not None:
            return True
        messagebox.showerror("缺少 Excel 组件", "当前 Python 环境未安装 openpyxl，无法读写 .xlsx 文件。请安装：python -m pip install openpyxl")
        return False

    def export_excel(self):
        if not self._has_openpyxl():
            return
        target = filedialog.asksaveasfilename(
            title="导出岗位记录",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
            initialfile="秋招岗位记录.xlsx",
        )
        if not target:
            return
        try:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "岗位记录"
            sheet.append([header for _, header in EXCEL_COLUMNS])
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            rows = self.conn.execute("SELECT * FROM jobs ORDER BY updated_at DESC, id DESC").fetchall()
            for row in rows:
                values = []
                for field, _header in EXCEL_COLUMNS:
                    raw = row[field]
                    if field in ("applied_date", "deadline"):
                        try:
                            value = self._excel_date_value(raw, with_time=False)
                        except ValueError:
                            value = self._cell_to_text(raw)
                    elif field == "interview_time":
                        try:
                            value = self._excel_date_value(raw, with_time=True)
                        except ValueError:
                            value = self._cell_to_text(raw)
                    else:
                        value = self._cell_to_text(raw)
                    values.append(value)
                sheet.append(values)
                for column, (field, _header) in enumerate(EXCEL_COLUMNS, start=1):
                    cell = sheet.cell(sheet.max_row, column)
                    if field in ("applied_date", "deadline") and isinstance(cell.value, date):
                        cell.number_format = "yyyy-mm-dd"
                    elif field == "interview_time" and isinstance(cell.value, datetime):
                        cell.number_format = "yyyy-mm-dd hh:mm"
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            widths = {"company": 18, "position": 24, "location": 14, "salary": 16, "channel": 14, "status": 12, "applied_date": 14, "interview_time": 22, "deadline": 14, "link": 32, "notes": 36}
            for column, (field, _header) in enumerate(EXCEL_COLUMNS, start=1):
                sheet.column_dimensions[sheet.cell(1, column).column_letter].width = widths.get(field, 16)
            workbook.save(target)
        except Exception as error:
            messagebox.showerror("导出失败", f"无法写入 Excel：{error}")
            return
        messagebox.showinfo("导出完成", f"已导出 {len(rows)} 条岗位记录。")

    def import_excel(self):
        if not self._has_openpyxl():
            return
        source = filedialog.askopenfilename(
            title="导入岗位记录",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")],
        )
        if not source:
            return
        workbook = None
        inserted = duplicates = invalid = 0
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
            sheet = workbook.active
            header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not header_row:
                raise ValueError("Excel 文件没有表头")
            mapping = {}
            for index, value in enumerate(header_row):
                header = self._cell_to_text(value)
                canonical = EXCEL_HEADER_ALIASES.get(header) or EXCEL_HEADER_ALIASES.get(header.lower())
                if canonical and canonical not in mapping:
                    mapping[canonical] = index
            if "company" not in mapping or "position" not in mapping:
                raise ValueError("表头至少需要包含“公司名称”和“岗位名称”")

            for row in sheet.iter_rows(min_row=2, values_only=True):
                def get_value(field):
                    index = mapping.get(field)
                    return self._cell_to_text(row[index], date_only=field in ("deadline", "applied_date")) if index is not None and index < len(row) else ""

                company, position = get_value("company"), get_value("position")
                if not company or not position:
                    invalid += 1
                    continue
                raw_values = {
                    "company": company,
                    "position": position,
                    "location": get_value("location"),
                    "salary": get_value("salary"),
                    "channel": get_value("channel"),
                    "status": get_value("status") or "待投递",
                    "deadline": get_value("deadline"),
                    "applied_date": get_value("applied_date"),
                    "interview_time": get_value("interview_time"),
                    "link": get_value("link"),
                    "notes": get_value("notes"),
                }
                try:
                    raw_values["deadline"] = self._normalize_date_value(raw_values["deadline"], with_time=False)
                    raw_values["applied_date"] = self._normalize_date_value(raw_values["applied_date"], with_time=False)
                    raw_values["interview_time"] = self._normalize_date_value(raw_values["interview_time"], with_time=True)
                except ValueError:
                    invalid += 1
                    continue
                if raw_values["status"] not in STATUSES:
                    raw_values["status"] = "待投递"
                duplicate = self.conn.execute(
                    "SELECT 1 FROM jobs WHERE company=? AND position=? AND applied_date=? AND COALESCE(link, '')=?",
                    (company, position, raw_values["applied_date"], raw_values["link"]),
                ).fetchone()
                if duplicate:
                    duplicates += 1
                    continue
                cursor = self.conn.execute(
                    "INSERT INTO jobs (company, position, location, salary, channel, status, deadline, applied_date, interview_time, link, notes) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    tuple(raw_values[field] for field in ("company", "position", "location", "salary", "channel", "status", "deadline", "applied_date", "interview_time", "link", "notes")),
                )
                inserted += 1
            self.conn.commit()
        except Exception as error:
            self.conn.rollback()
            messagebox.showerror("导入失败", f"无法读取 Excel：{error}")
            return
        finally:
            if workbook is not None:
                workbook.close()
        self.refresh()
        messagebox.showinfo("导入完成", f"新增 {inserted} 条，跳过重复 {duplicates} 条，无效 {invalid} 条。")

    def refresh_documents(self):
        if not hasattr(self, "document_tree"):
            return
        for item in self.document_tree.get_children():
            self.document_tree.delete(item)
        rows = self.conn.execute("SELECT * FROM documents ORDER BY created_at DESC, id DESC").fetchall()
        self.document_count_var.set(f"{len(rows)} 个文件")
        for row in rows:
            try:
                available = self._document_path(row).is_file()
            except (OSError, ValueError):
                available = False
            self.document_tree.insert(
                "", "end", iid=str(row["id"]),
                tags=("ready" if available else "missing",),
                values=(row["category"], row["original_name"], row["created_at"] or "—", "可访问" if available else "文件缺失"),
            )

    def _copy_document(self, category, source, source_attachment_id=None):
        source = Path(source).expanduser()
        if not source.is_file():
            raise FileNotFoundError(source)
        category = category if category in DOCUMENT_CATEGORIES else "其他"
        original_name = source.name
        stored_name = f"{uuid.uuid4().hex}_{original_name}"
        target = DOCUMENTS_DIR / stored_name
        shutil.copy2(source, target)
        try:
            cursor = self.conn.execute(
                "INSERT INTO documents (category, original_name, stored_path, source_attachment_id) VALUES (?, ?, ?, ?)",
                (category, original_name, stored_name, source_attachment_id),
            )
            self.conn.commit()
        except Exception:
            try:
                target.unlink()
            except OSError:
                pass
            raise
        return cursor.lastrowid

    def upload_documents(self, category=None):
        category = category or self.document_category_var.get()
        sources = filedialog.askopenfilenames(
            title=f"选择{category}文件",
            filetypes=[
                ("常用材料", "*.pdf *.doc *.docx *.xls *.xlsx *.png *.jpg *.jpeg *.txt"),
                ("所有文件", "*.*"),
            ],
        )
        if not sources:
            return
        success = failed = 0
        errors = []
        for source in sources:
            try:
                self._copy_document(category, source)
                success += 1
            except OSError as error:
                failed += 1
                errors.append(f"{Path(source).name}: {error}")
        self.refresh_documents()
        message = f"已复制 {success} 个文件。"
        if failed:
            message += f"\n失败 {failed} 个：\n" + "\n".join(errors[:5])
        messagebox.showinfo("材料上传完成", message)

    def _selected_document(self):
        selection = self.document_tree.selection()
        if not selection:
            return None
        return self.conn.execute("SELECT * FROM documents WHERE id = ?", (int(selection[0]),)).fetchone()

    def _document_path(self, row):
        root = DOCUMENTS_DIR.resolve()
        path = (root / row["stored_path"]).resolve()
        if root not in path.parents:
            raise ValueError("材料路径无效")
        return path

    def open_document(self):
        row = self._selected_document()
        if not row:
            return
        try:
            path = self._document_path(row)
            if not path.is_file():
                raise FileNotFoundError(path)
            if hasattr(os, "startfile"):
                os.startfile(str(path))
            else:
                webbrowser.open(path.as_uri())
        except (OSError, ValueError) as error:
            messagebox.showerror("打开材料失败", f"找不到或无法打开文件：{error}")

    def open_documents_folder(self):
        DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
        try:
            if hasattr(os, "startfile"):
                os.startfile(str(DOCUMENTS_DIR))
            else:
                webbrowser.open(DOCUMENTS_DIR.as_uri())
        except OSError as error:
            messagebox.showerror("打开目录失败", f"无法打开材料目录：{error}")

    def delete_document(self):
        row = self._selected_document()
        if not row or not messagebox.askyesno("确认删除", f"确定删除材料“{row['original_name']}”吗？"):
            return
        try:
            path = self._document_path(row)
            if path.is_file():
                path.unlink()
        except (OSError, ValueError) as error:
            messagebox.showerror("删除材料失败", f"无法删除文件：{error}")
            return
        self.conn.execute("DELETE FROM documents WHERE id = ?", (row["id"],))
        self.conn.commit()
        self.refresh_documents()

    @staticmethod
    def _ics_escape(value):
        return str(value or "").replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\r", "").replace("\n", "\\n")

    def export_calendar(self):
        target = filedialog.asksaveasfilename(
            title="导出日历文件",
            defaultextension=".ics",
            filetypes=[("iCalendar 文件", "*.ics")],
            initialfile="秋招日程.ics",
        )
        if not target:
            return
        rows = self.conn.execute("SELECT * FROM jobs ORDER BY deadline, interview_time, id").fetchall()
        stamp = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
        lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Campus Job Tracker//CN", "CALSCALE:GREGORIAN", "METHOD:PUBLISH"]
        event_count = 0
        for row in rows:
            if row["status"] in CALENDAR_EXCLUDED_STATUSES:
                continue
            summary_base = f"{row['company']} - {row['position']}"
            description = self._ics_escape("；".join(filter(None, (row["status"], row["notes"], row["link"]))))
            deadline = self._date_from_value(row["deadline"])
            if deadline:
                lines.extend([
                    "BEGIN:VEVENT",
                    f"UID:job-{row['id']}-deadline@campus-job-tracker",
                    f"DTSTAMP:{stamp}",
                    f"DTSTART;VALUE=DATE:{deadline.strftime('%Y%m%d')}",
                    f"DTEND;VALUE=DATE:{(deadline + timedelta(days=1)).strftime('%Y%m%d')}",
                    f"SUMMARY:{self._ics_escape('截止：' + summary_base)}",
                    f"DESCRIPTION:{description}",
                    "END:VEVENT",
                ])
                event_count += 1
            interview = row["interview_time"] or ""
            try:
                normalized_interview = self._normalize_date_value(interview, with_time=True) if interview else ""
            except ValueError:
                normalized_interview = ""
            if normalized_interview:
                if len(normalized_interview) > 10:
                    start = datetime.strptime(normalized_interview, "%Y-%m-%d %H:%M")
                    end = start + timedelta(hours=1)
                    time_lines = [
                        f"DTSTART:{start.strftime('%Y%m%dT%H%M%S')}",
                        f"DTEND:{end.strftime('%Y%m%dT%H%M%S')}",
                    ]
                else:
                    interview_date = date.fromisoformat(normalized_interview)
                    time_lines = [
                        f"DTSTART;VALUE=DATE:{interview_date.strftime('%Y%m%d')}",
                        f"DTEND;VALUE=DATE:{(interview_date + timedelta(days=1)).strftime('%Y%m%d')}",
                    ]
                lines.extend([
                    "BEGIN:VEVENT",
                    f"UID:job-{row['id']}-interview@campus-job-tracker",
                    f"DTSTAMP:{stamp}",
                    *time_lines,
                    f"SUMMARY:{self._ics_escape('笔试/面试：' + summary_base)}",
                    f"DESCRIPTION:{description}",
                    "END:VEVENT",
                ])
                event_count += 1
        lines.append("END:VCALENDAR")
        try:
            with Path(target).open("w", encoding="utf-8", newline="") as calendar_file:
                calendar_file.write("\r\n".join(lines) + "\r\n")
        except OSError as error:
            messagebox.showerror("日历导出失败", f"无法写入日历文件：{error}")
            return
        messagebox.showinfo("日历导出完成", f"已生成 {event_count} 个日历事件，可导入系统日历。")

    def check_deadline_reminders(self, show_empty=False):
        today = date.today()
        latest = today + timedelta(days=REMINDER_DAYS)
        reminders = []
        rows = self.conn.execute("SELECT company, position, status, deadline FROM jobs WHERE COALESCE(deadline, '') <> '' ORDER BY deadline").fetchall()
        for row in rows:
            if row["status"] in CALENDAR_EXCLUDED_STATUSES:
                continue
            due = self._date_from_value(row["deadline"])
            if not due or due > latest:
                continue
            days = (due - today).days
            label = "已逾期" if days < 0 else "今天到期" if days == 0 else f"{days} 天后到期"
            reminders.append(f"{row['company']} · {row['position']}（{label}，{due.isoformat()}）")
        if reminders:
            messagebox.showwarning("截止日期提醒", "以下岗位需要关注：\n\n" + "\n".join(reminders[:15]) + ("\n……" if len(reminders) > 15 else ""))
        elif show_empty:
            messagebox.showinfo("截止日期提醒", f"未来 {REMINDER_DAYS} 天内没有待处理的截止日期。")

    def query(self, sql, params=()):
        cur = self.conn.execute(sql, params)
        self.conn.commit()
        return cur

    def refresh(self):
        self.update_cards()
        self.load_jobs()

    def update_cards(self):
        rows = self.conn.execute("SELECT status, COUNT(*) AS n FROM jobs GROUP BY status").fetchall()
        counts = {row["status"]: row["n"] for row in rows}
        self.card_labels["全部"].configure(text=str(sum(counts.values())))
        self.card_labels["待投递"].configure(text=str(counts.get("待投递", 0)))
        ongoing = sum(counts.get(s, 0) for s in ONGOING_STATUSES)
        self.card_labels["进行中"].configure(text=str(ongoing))
        self.card_labels["Offer"].configure(text=str(counts.get("Offer", 0)))

    def load_jobs(self):
        if not hasattr(self, "tree"):
            return
        for item in self.tree.get_children():
            self.tree.delete(item)
        conditions, params = [], []
        keyword = self.search_var.get().strip()
        if keyword:
            conditions.append("(company LIKE ? OR position LIKE ? OR location LIKE ? OR salary LIKE ? OR channel LIKE ?)")
            params.extend([f"%{keyword}%"] * 5)
        selected = self.filter_var.get()
        if selected != "全部状态":
            if selected == "已终止":
                placeholders = ", ".join("?" for _ in TERMINATED_STATUSES)
                conditions.append(f"status IN ({placeholders})")
                params.extend(sorted(TERMINATED_STATUSES))
            else:
                conditions.append("status = ?")
                params.append(selected)
        where = " WHERE " + " AND ".join(conditions) if conditions else ""
        rows = self.conn.execute(f"SELECT * FROM jobs{where} ORDER BY CASE status WHEN '待投递' THEN 0 ELSE 1 END, updated_at DESC", params).fetchall()
        self.list_count_var.set(f"{len(rows)} 条记录")
        for row in rows:
            self.tree.insert("", "end", iid=str(row["id"]), tags=(row["status"],), values=(row["company"], row["position"], row["location"] or "—", row["salary"] or "—", row["channel"] or "—", row["status"], row["applied_date"] or "—", row["interview_time"] or "—", row["notes"] or "—", "已填写" if row["link"] else "—"))

    def select_job(self, _event=None):
        selection = self.tree.selection()
        if not selection:
            return
        self.selected_id = int(selection[0])
        row = self.conn.execute("SELECT * FROM jobs WHERE id = ?", (self.selected_id,)).fetchone()
        if not row:
            return
        for name, var in self.vars.items():
            var.set(row[name] or "")
        self.status_var.set(row["status"])
        self.notes.delete("1.0", "end")
        self.notes.insert("1.0", row["notes"] or "")
        self.save_button.configure(text="更新岗位")
        self.delete_button.configure(state="normal")
        self.open_button.configure(state="normal" if row["link"] else "disabled")

    def new_job(self):
        self.selected_id = None
        for var in self.vars.values():
            var.set("")
        self.status_var.set("待投递")
        self.notes.delete("1.0", "end")
        self.save_button.configure(text="保存岗位")
        self.delete_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.tree.selection_remove(self.tree.selection())

    def save_job(self):
        company, position = self.vars["company"].get().strip(), self.vars["position"].get().strip()
        if not company or not position:
            messagebox.showwarning("请补全信息", "公司名称和岗位名称为必填项。")
            return
        raw_values = {
            "company": company,
            "position": position,
            "location": self.vars["location"].get().strip(),
            "salary": self.vars["salary"].get().strip(),
            "channel": self.vars["channel"].get().strip(),
            "status": self.status_var.get(),
            "deadline": self.vars["deadline"].get().strip(),
            "applied_date": self.vars["applied_date"].get().strip(),
            "interview_time": self.vars["interview_time"].get().strip(),
            "link": self.vars["link"].get().strip(),
            "notes": self.notes.get("1.0", "end").strip(),
        }
        normalized = self._normalize_job_dates(raw_values)
        if normalized is None:
            return
        values = tuple(normalized[field] for field in ("company", "position", "location", "salary", "channel", "status", "deadline", "applied_date", "interview_time", "link", "notes"))
        if self.selected_id is None:
            cur = self.query("INSERT INTO jobs (company, position, location, salary, channel, status, deadline, applied_date, interview_time, link, notes) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", values)
            self.selected_id = cur.lastrowid
            messagebox.showinfo("已保存", "岗位已加入你的秋招清单。")
        else:
            self.query("UPDATE jobs SET company=?, position=?, location=?, salary=?, channel=?, status=?, deadline=?, applied_date=?, interview_time=?, link=?, notes=?, updated_at=CURRENT_TIMESTAMP WHERE id=?", values + (self.selected_id,))
            messagebox.showinfo("已更新", "岗位信息已更新。")
        self.refresh()
        self.tree.selection_set(str(self.selected_id))

    def delete_job(self):
        if self.selected_id is None:
            return
        if messagebox.askyesno("确认删除", "确定要删除这条岗位记录吗？"):
            self.query("DELETE FROM jobs WHERE id = ?", (self.selected_id,))
            self.new_job()
            self.refresh()

    def open_link(self):
        link = self.vars["link"].get().strip()
        if link:
            webbrowser.open(link if link.startswith(("http://", "https://")) else "https://" + link)

    def close(self):
        self.conn.close()
        self.destroy()


if __name__ == "__main__":
    if sys.platform.startswith("win"):
        try:
            from ctypes import windll
            windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    JobTracker().mainloop()
